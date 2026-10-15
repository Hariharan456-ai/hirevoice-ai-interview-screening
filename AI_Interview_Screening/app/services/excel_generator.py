import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .. import config, database as db
from .ranking import ranking_table

COLS = ["Rank", "Candidate Name", "Phone", "Email", "College", "Degree", "Specialization", "Graduation Year",
        "Questions Completed", "Total Score", "Maximum Score", "Percentage", "Avg Response Time (s)", "AI Suspected", "Report PDF (click)"]

def generate_excel():
    wb = Workbook(); first = True
    for role in db.query("SELECT * FROM job_roles ORDER BY name"):
        rows = [r for r in ranking_table(role["id"]) if r["rank"] is not None]
        if not rows: continue
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = re.sub(r"[\\/*?:\[\]]", "", role["name"])[:31]
        ws.append(COLS)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0B1F44"); c.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append([r["rank"], r["name"], r["phone"], r["email"], r["college"], r["degree"], r["specialization"],
                       r["graduation_year"], r["done_count"], r["total_score"], r["max_score"], (r["percentage"] or 0) / 100,
                       round(r["avg_response_time"], 1) if r["avg_response_time"] is not None else "",
                       "YES - verify" if r["ai_flag"] else "No", ""])
            row = ws.max_row; ws.cell(row, 12).number_format = "0.0%"
            if r["ai_flag"]:
                for col in (2, 14): ws.cell(row, col).font = Font(bold=True, color="DC2626")
                ws.cell(row, 14).fill = PatternFill("solid", fgColor="FEE2E2")
            link = r["drive_link"] or (("file:///" + str(r["pdf_path"]).replace("\\", "/")) if r["pdf_path"] else "")
            if link:
                cell = ws.cell(row, 15, "Open Q&A PDF (Drive)" if r["drive_link"] else "Open Q&A PDF"); cell.hyperlink = link
                cell.font = Font(color="1D4ED8", underline="single")
        for i, w in enumerate([7, 26, 16, 28, 30, 12, 24, 10, 12, 11, 11, 11, 14, 14, 18], 1):
            ws.column_dimensions[chr(64 + i)].width = w
        ws.freeze_panes = "A2"
    if first:
        wb.active.title = "Ranking"; wb.active.append(COLS)
    out = config.EXPORT_FOLDER / "candidate_ranking.xlsx"
    wb.save(out)
    return out
